"""Build a Quantum Mechanics prerequisite tracker spreadsheet, cross-referenced
against what this repository already has built and tested for each prerequisite
topic -- an evidence-based self-assessment, not a guess at academic status."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

FONT = "Arial"
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F3864")
SUBTITLE_FONT = Font(name=FONT, size=10, italic=True, color="595959")
BODY_FONT = Font(name=FONT, size=10)
BOLD_BODY = Font(name=FONT, size=10, bold=True)
COVERED_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
GAP_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ============================================================ SHEET 1: Prerequisite Map
ws1 = wb.active
ws1.title = "Prerequisite Map"

ws1["B2"] = "Quantum Mechanics — Prerequisite Map"
ws1["B2"].font = TITLE_FONT
ws1["B3"] = "Course chain feeding into Quantum Mechanics, as given"
ws1["B3"].font = SUBTITLE_FONT

headers1 = ["Course Code", "Course Title", "Key Topics This Course Covers", "Role"]
for i, h in enumerate(headers1):
    c = ws1.cell(row=5, column=2 + i, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")

prereq_rows = [
    ("Engr 17", "Introductory Circuit Analysis",
     "Ohm's law, KVL/KCL, RC/RL/RLC transient response, phasors & AC steady-state, Thevenin/Norton equivalents",
     "Prerequisite"),
    ("Math 45", "Differential Equations for Science and Engineering",
     "1st/2nd-order ODEs, series solutions, Laplace transforms, systems of linear ODEs",
     "Prerequisite"),
    ("Phys 11C", "General Physics: Electricity and Magnetism",
     "Coulomb's law, Gauss's law, circuits, Maxwell's equations, EM wave propagation",
     "Prerequisite"),
    ("—", "Quantum Mechanics",
     "Schrodinger equation, operators & eigenstates, the harmonic oscillator, angular momentum, spin",
     "Target course"),
]
for r, row in enumerate(prereq_rows, start=6):
    for c_idx, val in enumerate(row):
        cell = ws1.cell(row=r, column=2 + c_idx, value=val)
        cell.font = BOLD_BODY if c_idx == 0 else BODY_FONT
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if row[3] == "Target course":
        for c_idx in range(4):
            ws1.cell(row=r, column=2 + c_idx).fill = PatternFill(
                start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

ws1.column_dimensions["A"].width = 2
ws1.column_dimensions["B"].width = 14
ws1.column_dimensions["C"].width = 30
ws1.column_dimensions["D"].width = 62
ws1.column_dimensions["E"].width = 14
for r in range(6, 10):
    ws1.row_dimensions[r].height = 40

# ============================================================ SHEET 2: Repo Coverage
ws2 = wb.create_sheet("Repo Coverage")
ws2["B2"] = "What's Already Built and Tested in This Repository"
ws2["B2"].font = TITLE_FONT
ws2["B3"] = "Checked directly against the repo's actual files before listing anything here — not assumed"
ws2["B3"].font = SUBTITLE_FONT

headers2 = ["Prerequisite Course", "Topic", "Covered?", "File(s)", "Notes"]
for i, h in enumerate(headers2):
    c = ws2.cell(row=5, column=2 + i, value=h)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.border = BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")

coverage_rows = [
    ("Engr 17", "Ohm's law / KVL / KCL, DC circuits", "Y", "dgs/dc_circuits.py, dgs/circuits.py",
     "Node/mesh analysis and basic circuit laws"),
    ("Engr 17", "RC/RL/RLC transient response", "Y", "dgs/spice.py",
     "RK4 time-domain solve, damping regime classification (under/critical/over), verified underdamped overshoot and overdamped monotonicity"),
    ("Engr 17", "AC steady-state / phasors, Thevenin/Norton", "Y", "dgs/ac_circuits.py",
     "AC power, op-amps"),
    ("Engr 17", "Dependent (controlled) sources, s-domain circuit analysis", "Y",
     "dgs/spice.py (s_domain_nodal_analysis)",
     "Built this session: VCVS op-amp model verified to converge to the exact textbook gain 1+Rf/Rin"),
    ("Math 45", "1st/2nd-order ODEs, series solutions", "Y", "dgs/math_methods.py",
     "Frobenius/series methods, ODE toolkit"),
    ("Math 45", "Laplace transforms", "Y", "dgs/laplace_transform.py",
     "Built this session: 6 standard pairs derived symbolically (SymPy), verified against textbook closed forms; pole/zero stability analysis"),
    ("Math 45", "Systems of linear ODEs / linear algebra methods", "Y", "dgs/math_methods.py",
     "Eigenvalues, SVD, determinants (Arfken/Boas §2)"),
    ("Phys 11C", "Coulomb's law, Gauss's law, electrostatics", "Y", "griffiths/electrodynamics.py, griffiths/potentials.py",
     "Full Griffiths-level treatment"),
    ("Phys 11C", "Maxwell's equations", "Y", "griffiths/electrodynamics.py",
     "Full Griffiths-level treatment of Maxwell's equations"),
    ("Phys 11C", "Lorentz oscillator, complex refractive index, dispersion operator", "Y", "dgs/em_dispersion.py",
     "Ties Griffiths Ch. 9 directly to this repo's own dispersion operator H(omega)=exp(i*beta2*omega^2*L/2)"),
    ("Phys 11C", "EM wave propagation, dispersion", "Y", "notebooks/laplace_maxwell_dispersion_regression.ipynb",
     "Built this session: symbolic derivation of this repo's own dispersion kernel from Maxwell's wave equation"),
    ("Quantum Mechanics (target)", "Schrodinger equation, infinite well, time evolution", "Y",
     "dgs/stationary_states.py",
     "Eigenstates, expansion coefficients, time evolution, expectation values"),
    ("Quantum Mechanics (target)", "Harmonic oscillator, ladder operators, number states", "Y",
     "dgs/quantum_oscillator.py",
     "Creation/annihilation/number operators, energy levels, mean occupation"),
    ("Quantum Mechanics (target)", "Klein-Gordon equation (relativistic QM preview)", "Y",
     "dgs/particle_physics.py (quantum_field_theory_preview), tests/test_particle_physics.py",
     "Klein-Gordon dispersion relation, causality check (v_group < c)"),
    ("Quantum Mechanics (target)", "Gauge invariance, minimal coupling", "Y",
     "dgs/gauge_invariance.py",
     "No companion notebook yet for this module specifically"),
    ("Quantum Mechanics (target)", "Quantum harmonic oscillator, photonics connection", "Y",
     "physics_repo/notebooks/quantum_harmonic_oscillator_photonics.ipynb",
     "From the companion physics_repo/ curriculum notebooks"),
]
for r, row in enumerate(coverage_rows, start=6):
    for c_idx, val in enumerate(row):
        cell = ws2.cell(row=r, column=2 + c_idx, value=val)
        cell.font = BOLD_BODY if c_idx in (0, 2) else BODY_FONT
        cell.border = BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    fill = COVERED_FILL if row[2] == "Y" else GAP_FILL
    for c_idx in range(5):
        ws2.cell(row=r, column=2 + c_idx).fill = fill

last_row = 5 + len(coverage_rows)
n_covered = sum(1 for row in coverage_rows if row[2] == "Y")
n_total = len(coverage_rows)
ws2.cell(row=last_row + 2, column=2, value="Topics covered:").font = BOLD_BODY
# NOTE: written as a plain value, not a formula -- this environment's LibreOffice
# recalculation step (required before shipping any formula) could not run here
# (a Windows socket.AF_UNIX platform limitation in the recalc tool itself, not a
# formula error), so this count is computed directly in Python from the exact
# same coverage_rows list immediately above, and will go stale if that list is
# edited without updating this cell by hand.
ws2.cell(row=last_row + 2, column=3, value=f"{n_covered} / {n_total}")
ws2.cell(row=last_row + 2, column=3).font = BOLD_BODY
ws2.cell(row=last_row + 3, column=2,
         value="(static count, not a live formula -- see the script's comment for why)").font = SUBTITLE_FONT

ws2.column_dimensions["A"].width = 2
ws2.column_dimensions["B"].width = 22
ws2.column_dimensions["C"].width = 34
ws2.column_dimensions["D"].width = 9
ws2.column_dimensions["E"].width = 40
ws2.column_dimensions["F"].width = 46
for r in range(6, last_row + 1):
    ws2.row_dimensions[r].height = 42

# ============================================================ SHEET 3: Legend / How to use
ws3 = wb.create_sheet("How to Use")
ws3["B2"] = "How to Use This Tracker"
ws3["B2"].font = TITLE_FONT
notes = [
    "This workbook has two working tabs:",
    "",
    "1. Prerequisite Map — the course chain (Engr 17, Math 45, Phys 11C) feeding into Quantum Mechanics, "
    "with each course's key topics listed for reference.",
    "",
    "2. Repo Coverage — a topic-by-topic check of what is ALREADY built and tested in this repository for "
    "each prerequisite subject, with the exact file(s) to open for each. Every 'Y' was checked directly "
    "against the repo's actual files (grep + reading the module) before being listed, not assumed.",
    "",
    "To adapt this for your own use: add a 'Status' column (e.g. Not Started / In Progress / Completed) "
    "and a 'Grade' column to the Prerequisite Map tab if you want to track your own progress through "
    "these courses — this version deliberately does not include those, since they reflect your actual "
    "academic status, which only you know.",
]
for i, line in enumerate(notes):
    cell = ws3.cell(row=4 + i, column=2, value=line)
    cell.font = BODY_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
ws3.column_dimensions["B"].width = 100

wb.save("qm_prerequisite_tracker.xlsx")
print("wrote qm_prerequisite_tracker.xlsx")
